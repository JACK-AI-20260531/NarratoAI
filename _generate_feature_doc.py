from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from pathlib import Path

output = Path(r"d:\work\ai\NarratoAI\JackAI视频脚本功能说明书.docx")

doc = Document()
styles = doc.styles
styles['Normal'].font.name = 'Microsoft YaHei'
styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
styles['Normal'].font.size = Pt(10.5)
for style_name in ['Title', 'Heading 1', 'Heading 2', 'Heading 3']:
    style = styles[style_name]
    style.font.name = 'Microsoft YaHei'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')


def add_title(text):
    p = doc.add_heading(text, 0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_h1(text):
    doc.add_heading(text, level=1)


def add_h2(text):
    doc.add_heading(text, level=2)


def add_p(text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_bullets(items):
    for item in items:
        doc.add_paragraph(item, style='List Bullet')


def add_numbered(items):
    for item in items:
        doc.add_paragraph(item, style='List Number')


def add_table(headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph('')


add_title('JackAI 视频脚本功能说明书')
add_p('本文档用于说明 JackAI 中“视频类型”模块的六个功能入口，帮助用户根据素材情况选择合适的脚本生成方式。')

add_h1('一、功能概述')
add_p('JackAI 的“视频类型”模块用于生成或加载视频脚本。视频脚本是后续生成成片的核心数据，包含每个片段的时间轴、画面描述、解说词、是否保留原声等信息。')
add_p('当前支持六种功能模式：选择/上传脚本、逐帧解说、短剧混剪、短剧解说、电影原片解说、电影混剪。')

sections = [
    ('二、选择/上传脚本', '用于加载已经存在的视频脚本文件。该功能不调用 AI 生成新脚本，而是直接使用已有 JSON 脚本。',
     ['已经生成过脚本，希望再次生成视频。', '手动修改过脚本，希望使用修改后的版本生成视频。', '希望更换配音、字幕样式、视频比例，但不想重新生成脚本。', '有外部导入的标准 JSON 脚本。'],
     ['脚本文件，格式为 .json。', '原始视频文件。'],
     '将脚本加载到“视频脚本”编辑框，用户可以继续编辑脚本，或直接点击“生成视频”。',
     ['选择“选择/上传脚本”。', '从脚本文件列表中选择历史脚本，或上传新的 JSON 文件。', '选择对应的视频文件。', '检查脚本内容。', '点击“生成视频”。'],
     '如果脚本已经生成过，优先使用该功能，可以节省时间和模型费用。'),
    ('三、逐帧解说', '用于让 AI 根据视频画面自动生成解说脚本。系统会从视频中按固定间隔抽取关键帧，调用视觉模型分析画面，再生成对应的视频解说脚本。',
     ['用户只有视频，没有文案。', '视频没有字幕，无法根据字幕生成脚本。', '用户希望 AI 根据画面内容自动写解说。', '适合纪录片、旅行视频、产品演示、实拍素材、无字幕素材等。'],
     ['原始视频文件。', '视频主题。', '自定义生成提示词。', '抽帧间隔。', '视觉模型批处理数量。'],
     '自动生成带时间轴的视频脚本 JSON。',
     ['选择“逐帧解说”。', '选择视频文件。', '输入视频主题。', '可选填写生成提示词。', '设置抽帧间隔和批处理数量。', '点击“生成视频脚本”。', '检查并保存脚本。', '点击“生成视频”。'],
     '抽帧间隔越小，画面分析越细，但消耗更多 token，生成速度也会更慢。普通视频可设置为 3 秒左右，长视频建议适当调大。'),
    ('四、短剧混剪', '用于从短剧或剧情视频中自动生成多个精彩片段脚本。该功能更偏向“切片”和“混剪”，不是完整剧情解说。',
     ['有一集短剧，想快速剪出多个片段。', '想做短视频矩阵切片。', '想从长视频中提取多个精彩片段。', '不想手动拖时间轴找爆点。'],
     ['短剧视频文件。', '需要生成的片段数量。'],
     '生成多个短视频片段脚本，每个片段包含对应的视频时间范围、画面描述、旁白或原声设置。',
     ['选择“短剧混剪”。', '选择短剧视频文件。', '设置需要生成的片段数量。', '点击“生成短视频脚本”。', '检查脚本内容。', '保存脚本。', '点击“生成视频”。'],
     '适合做爆点切片、剧情反转片段、人物冲突片段。如果目标是完整讲完一集剧情，更建议使用“短剧解说”。'),
    ('五、短剧解说', '用于根据短剧字幕生成剧情解说脚本。该功能主要依赖字幕内容，而不是逐帧分析画面。',
     ['有短剧字幕文件。', '想将短剧对白整理成解说文案。', '想做“几分钟看完一集短剧”类型视频。', '没有字幕时，可以通过 Fun-ASR 自动转写字幕。'],
     ['短剧视频文件。', '字幕文件 .srt。', '或上传音视频，通过阿里百炼 Fun-ASR 自动转写字幕。', '短剧名称。', 'temperature 参数。'],
     '生成短剧剧情解说脚本 JSON。',
     ['选择“短剧解说”。', '选择视频文件。', '上传 .srt 字幕文件。', '如果没有字幕，可使用 Fun-ASR 转写。', '输入短剧名称。', '设置 temperature。', '点击“生成短剧解说脚本”。', '检查并保存脚本。', '点击“生成视频”。'],
     '短剧剧情主要依赖对白时，优先使用该功能。如果字幕质量差，建议先修正字幕。temperature 越高，文案越发散；越低，文案越稳定。'),
    ('六、电影原片解说', '用于根据已有电影解说文案，自动匹配原电影画面。适合“先有解说稿，再找原片画面”的工作流。',
     ['已经写好了电影解说文案。', '想自动从电影原片中找到对应画面。', '想制作完整电影解说视频。', '不想手动在电影中寻找每段文案对应的时间点。'],
     ['电影原片文件。', '电影原片解说文案。', '抽帧间隔。', '视觉模型批处理数量。', '候选画面数量。'],
     '生成匹配好原片画面的解说脚本 JSON，每段解说会对应电影中的具体时间范围。',
     ['选择“电影原片解说”。', '选择电影原片文件。', '输入或粘贴电影解说文案。', '设置抽帧间隔、批处理数量、候选画面数。', '点击“匹配原片画面”。', '等待匹配完成。', '检查脚本。', '保存脚本。', '点击“生成视频”。'],
     '适合完整电影解说。文案越清晰，画面匹配效果越好。电影越长，分析时间和视觉模型消耗越高。'),
    ('七、电影混剪', '用于根据主题或关键词，从电影原片中自动挑选相关镜头，生成混剪脚本。适合“按主题找画面”。',
     ['想做电影高光混剪。', '想做人物向、情绪向、主题向视频。', '只想找出某类镜头，例如反击、爱情、悬疑、动作、催泪片段。', '没有完整解说稿，只想根据关键词自动生成混剪内容。'],
     ['电影原片文件。', '混剪主题或关键词。', '抽帧间隔。', '视觉模型批处理数量。', '混剪片段数。', '单段时长。'],
     '生成电影混剪脚本 JSON，脚本包含多个与主题相关的电影片段。',
     ['选择“电影混剪”。', '选择电影原片文件。', '输入混剪主题或关键词。', '设置混剪片段数和单段时长。', '点击“生成电影混剪”。', '检查脚本。', '保存脚本。', '点击“生成视频”。'],
     '适合做短视频平台的电影高光合集。主题越明确，匹配效果越好，例如：男主高能反击、悬疑反转、爱情名场面、动作追逐、悲伤催泪片段。'),
]

for title, desc, scenes, inputs, output_text, steps, tips in sections:
    add_h1(title)
    add_h2('功能定位')
    add_p(desc)
    add_h2('适用场景')
    add_bullets(scenes)
    add_h2('输入内容')
    add_bullets(inputs)
    add_h2('输出结果')
    add_p(output_text)
    add_h2('典型使用流程')
    add_numbered(steps)
    add_h2('使用建议')
    add_p(tips)

add_h1('八、功能选择建议')
add_table(['目标', '推荐功能'], [
    ['已经有 JSON 脚本', '选择/上传脚本'],
    ['只有视频，没有文案和字幕', '逐帧解说'],
    ['想从短剧里切多个精彩片段', '短剧混剪'],
    ['有短剧字幕，想生成剧情解说', '短剧解说'],
    ['已有电影解说稿，想匹配原片画面', '电影原片解说'],
    ['没有电影解说稿，只想按主题剪电影镜头', '电影混剪'],
])

add_h1('九、简单决策流程')
add_p('如果已经有脚本，直接选择“选择/上传脚本”。')
add_p('如果没有脚本，先判断素材类型：')
add_bullets([
    '短剧素材：想做多个短视频切片，用“短剧混剪”；想做完整剧情解说，用“短剧解说”。',
    '电影素材：已经有解说文案，用“电影原片解说”；只有主题关键词，用“电影混剪”。',
    '普通视频素材：没有字幕也没有文案，用“逐帧解说”。',
])

add_h1('十、注意事项')
add_bullets([
    '生成脚本后，建议先检查“视频脚本”内容，再保存脚本。',
    '保存后的脚本可以在“选择/上传脚本”中重复使用。',
    '长视频、电影原片、逐帧分析类功能会消耗更多视觉模型 token。',
    '如果只是换声音、换字幕、换视频比例，不需要重新生成脚本。',
    '如果生成结果不理想，可以手动修改脚本后重新保存。',
])

doc.save(output)
print(output)
